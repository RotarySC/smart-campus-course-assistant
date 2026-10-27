"""Parser Adapter：txt / md / pdf（文本层）/ docx → 正常化文本 + 固定窗口分块。

设计依据：后端详细设计 §10.1、§10.3；SRS FR-DOC-002（空文本/损坏/加密要可解释失败）。
不承诺扫描件 OCR（PHASE-PLAN §3「明确不做」第 3 项）。
"""

from __future__ import annotations

import hashlib
import io
import os
from dataclasses import dataclass, field
from typing import Final

from app.shared.ids import PREFIX_CHUNK, new_id

CHUNK_SIZE_DEFAULT: Final = 500
CHUNK_OVERLAP_DEFAULT: Final = 80

SUPPORTED_TYPES: Final = ("txt", "md", "pdf", "docx")


class ParserError(Exception):
    """解析失败（确定性失败，不做自动重试）。code 落入 background_tasks.error_code。"""

    def __init__(self, code: str, message: str = "") -> None:
        self.code = code
        self.message = message or code
        super().__init__(self.message)


@dataclass(frozen=True)
class Segment:
    """结构段：pdf 一页 / docx 一段 / txt 与 md 全文一段。"""

    page_number: int | None
    paragraph_index: int | None
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    file_type: str
    segments: list[Segment]
    text: str = field(init=False)

    def __post_init__(self) -> None:
        object.__setattr__(self, "text", normalize_text("\n\n".join(s.text for s in self.segments)))


@dataclass(frozen=True)
class ChunkDraft:
    chunk_id: str
    chunk_index: int
    page_number: int | None
    location_label: str
    text: str
    content_digest: bytes


def chunk_size() -> int:
    return int(os.environ.get("CCA_CHUNK_SIZE", CHUNK_SIZE_DEFAULT))


def chunk_overlap() -> int:
    return int(os.environ.get("CCA_CHUNK_OVERLAP", CHUNK_OVERLAP_DEFAULT))


def parser_config_digest() -> bytes:
    # 摘要必须覆盖"会让向量不同源"的全部冻结参数：分块参数 + Embedding Provider/维度。
    # 只写不读时（旧实现），参数变更后旧版本仍能发布，检索查询向量与库内向量不同源，
    # 表现为"已发布但永久检索不到"。
    from app.adapters.embedding.registry import current_dim, provider_name

    raw = (
        f"parser-v1|chunk={chunk_size()}|overlap={chunk_overlap()}|norm=lf-strip"
        f"|embed={provider_name()}:{current_dim()}"
    )
    return hashlib.sha256(raw.encode("utf-8")).digest()


def normalize_text(text: str) -> str:
    """换行统一为 ``\\n``，去掉行尾空白，整体 trim；不改变正文与码点。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    return "\n".join(lines).strip()


# --------------------------------------------------------------------------- #
# 格式解析
# --------------------------------------------------------------------------- #
def detect_file_type(filename: str) -> str | None:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return suffix if suffix in SUPPORTED_TYPES else None


def parse_bytes(data: bytes, file_type: str) -> ParsedDocument:
    """按文件类型解析为结构段；失败抛 :class:`ParserError`。"""
    if file_type not in SUPPORTED_TYPES:
        raise ParserError("UNSUPPORTED_FILE_TYPE", "不支持的文件类型")
    if not data:
        raise ParserError("NO_TEXT_EXTRACTED", "文件为空")
    if file_type in ("txt", "md"):
        return _parse_plain(data, file_type)
    if file_type == "pdf":
        return _parse_pdf(data)
    return _parse_docx(data)


def _decode_plain(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise ParserError("PARSE_FAILED", "文本不是合法 UTF-8") from exc
    if "\x00" in text:
        raise ParserError("PARSE_FAILED", "文本包含二进制内容")
    return text


def _parse_plain(data: bytes, file_type: str) -> ParsedDocument:
    text = normalize_text(_decode_plain(data))
    if not text:
        raise ParserError("NO_TEXT_EXTRACTED", "文件没有可抽取的文本")
    return ParsedDocument(file_type=file_type, segments=[Segment(None, None, text)])


def _parse_pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ParserError("DOCUMENT_ENCRYPTED", "PDF 已加密，无法读取文本层")
            except ParserError:
                raise
            except Exception as exc:  # noqa: BLE001 - 加密 PDF 驱动差异
                raise ParserError("DOCUMENT_ENCRYPTED", "PDF 已加密，无法读取文本层") from exc
        segments: list[Segment] = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                page_text = normalize_text(page.extract_text() or "")
            except Exception as exc:  # noqa: BLE001 - 单页失败视为文档损坏
                raise ParserError("PARSE_FAILED", f"PDF 第 {index} 页解析失败") from exc
            if page_text:
                segments.append(Segment(page_number=index, paragraph_index=None, text=page_text))
    except ParserError:
        raise
    except (PdfReadError, ValueError, OSError) as exc:
        raise ParserError("PARSE_FAILED", "PDF 文件损坏或不是有效 PDF") from exc
    if not segments:
        raise ParserError("NO_TEXT_EXTRACTED", "PDF 没有文本层（可能为扫描件）")
    return ParsedDocument(file_type="pdf", segments=segments)


def _parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document as DocxDocument  # python-docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = DocxDocument(io.BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError, OSError) as exc:
        raise ParserError("PARSE_FAILED", "DOCX 文件损坏或不是有效 DOCX") from exc
    segments: list[Segment] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = normalize_text(paragraph.text or "")
        if text:
            segments.append(Segment(page_number=None, paragraph_index=index, text=text))
    if not segments:
        raise ParserError("NO_TEXT_EXTRACTED", "DOCX 没有可抽取的段落文本")
    return ParsedDocument(file_type="docx", segments=segments)


# --------------------------------------------------------------------------- #
# 分块
# --------------------------------------------------------------------------- #
def build_chunks(parsed: ParsedDocument) -> list[ChunkDraft]:
    """固定窗口 + 重叠；不跨结构段（pdf 逐页 / docx 逐段）拼接，保留页码映射。"""
    size = max(50, chunk_size())
    overlap = max(0, min(chunk_overlap(), size // 2))
    chunks: list[ChunkDraft] = []
    for segment in parsed.segments:
        text = segment.text
        if not text:
            continue
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            piece = text[start:end].strip()
            if piece:
                chunks.append(_make_chunk(len(chunks), segment, piece))
            if end >= len(text):
                break
            start = end - overlap
    if not chunks:
        raise ParserError("NO_TEXT_EXTRACTED", "分块后没有有效文本")
    return chunks


def _make_chunk(index: int, segment: Segment, text: str) -> ChunkDraft:
    if segment.page_number is not None:
        label = f"第 {segment.page_number} 页"
    elif segment.paragraph_index is not None:
        label = f"段落 {segment.paragraph_index}"
    else:
        label = f"片段 {index + 1}"
    return ChunkDraft(
        chunk_id=new_id(PREFIX_CHUNK),
        chunk_index=index,
        page_number=segment.page_number,
        location_label=label[:100],
        text=text,
        content_digest=hashlib.sha256(text.encode("utf-8")).digest(),
    )


def manifest_payload(parsed: ParsedDocument, chunks: list[ChunkDraft]) -> dict:
    """chunk-manifest 落盘内容（含 chunk_id/序号/页码/文本/长度）。"""
    return {
        "schema_version": 1,
        "file_type": parsed.file_type,
        "text_length": len(parsed.text),
        "chunk_size": chunk_size(),
        "chunk_overlap": chunk_overlap(),
        "chunk_count": len(chunks),
        "chunks": [
            {
                "chunk_id": c.chunk_id,
                "chunk_index": c.chunk_index,
                "page_number": c.page_number,
                "location_label": c.location_label,
                "length": len(c.text),
                "text": c.text,
            }
            for c in chunks
        ],
    }


__all__ = [
    "CHUNK_OVERLAP_DEFAULT",
    "CHUNK_SIZE_DEFAULT",
    "SUPPORTED_TYPES",
    "ChunkDraft",
    "ParsedDocument",
    "ParserError",
    "Segment",
    "build_chunks",
    "chunk_overlap",
    "chunk_size",
    "detect_file_type",
    "manifest_payload",
    "normalize_text",
    "parse_bytes",
    "parser_config_digest",
]
